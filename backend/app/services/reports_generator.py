"""Enterprise Report Generator Service for Asclepius — Generates styled Excel and PDF deliverables."""
import io
import os
import json
from datetime import datetime, UTC

from fpdf import FPDF
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from dotenv import load_dotenv
load_dotenv()

GROQ_MODEL = os.getenv("GROQ_MODEL", "llama-3.3-70b-versatile")


def get_ai_narrative(title, summary):
    try:
        from groq import Groq
        groq_key = os.environ.get("GROQ_API_KEY")
        if not groq_key:
            return None
            
        client = Groq(api_key=groq_key)
        
        # Filter context to avoid context length issues
        t_lower = title.lower()
        if "sales" in t_lower:
            context_data = summary.get("sales", {})
        elif "inventory" in t_lower:
            context_data = summary.get("inventory", {})
        elif "anomaly" in t_lower:
            context_data = {"anomalies": summary.get("anomalies", [])}
        else:
            context_data = summary.get("kpis", {})
            
        prompt = f"""
        You are Asclepius, an expert enterprise pharmaceutical data analyst.
        Write a very concise, professional 2-paragraph executive narrative for a report titled "{title}".
        Extract 2-3 key insights from the following JSON data.
        Use plain text only. Do NOT use markdown. Do NOT use asterisks. Do NOT include greetings.
        Keep it highly professional, data-driven, and actionable. Focus strictly on the data provided.
        
        DATA:
        {json.dumps(context_data)[:2500]}
        """
        
        resp = client.chat.completions.create(
            model=GROQ_MODEL,
            messages=[{"role": "user", "content": prompt}],
            temperature=0.3,
            max_tokens=300
        )
        return resp.choices[0].message.content.strip()
    except Exception as e:
        err_msg = str(e)
        if "model_decommissioned" in err_msg or "deprecated" in err_msg:
            print("Groq model deprecated, check https://console.groq.com/docs/deprecations")
        print("Groq Error:", e)
        return None


def generate_excel_report(title: str, dataset=None) -> bytes:
    """Generate a multi-sheet executive pharmaceutical analytics Excel workbook."""
    if not dataset or not dataset.analysis_summary:
        raise ValueError("Cannot generate report: No dataset or analysis summary provided.")
    
    summary = dataset.analysis_summary
    wb = Workbook()

    # Colors & Styles
    navy_fill = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid")
    indigo_fill = PatternFill(start_color="3730A3", end_color="3730A3", fill_type="solid")
    soft_indigo_fill = PatternFill(start_color="EEF2FF", end_color="EEF2FF", fill_type="solid")
    light_grey_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
    green_fill = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid")
    red_fill = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")
    amber_fill = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")

    title_font = Font(name="Calibri", size=16, bold=True, color="FFFFFF")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    subhead_font = Font(name="Calibri", size=12, bold=True, color="1E293B")
    body_font = Font(name="Calibri", size=11, color="334155")
    bold_font = Font(name="Calibri", size=11, bold=True, color="0F172A")
    green_font = Font(name="Calibri", size=11, bold=True, color="166534")
    amber_font = Font(name="Calibri", size=11, bold=True, color="92400E")

    thin_border = Border(
        left=Side(style="thin", color="CBD5E1"),
        right=Side(style="thin", color="CBD5E1"),
        top=Side(style="thin", color="CBD5E1"),
        bottom=Side(style="thin", color="CBD5E1"),
    )

    # Sheet 1: Executive Dashboard
    ws1 = wb.active
    ws1.title = "Executive KPI Summary"
    ws1.views.sheetView[0].showGridLines = True

    ws1.merge_cells("A1:C2")
    cell_top = ws1["A1"]
    cell_top.value = f"  ASCLEPIUS — {title.upper()}"
    cell_top.font = title_font
    cell_top.fill = navy_fill
    cell_top.alignment = Alignment(vertical="center")

    ws1["A4"] = "KEY OPERATIONAL METRICS & FINANCIAL HEALTH"
    ws1["A4"].font = subhead_font

    headers_kpi = ["Metric Name", "Current Value", "Status"]
    ws1.append([])
    ws1.append([])
    ws1.append(headers_kpi)

    for col_num, h_text in enumerate(headers_kpi, 1):
        cell = ws1.cell(row=5, column=col_num)
        cell.font = header_font
        cell.fill = indigo_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border

    ds_kpis = summary.get("kpis", {})
    anoms = summary.get("anomalies", [])
    health = dataset.business_health_score if dataset.business_health_score else "N/A"
    
    kpis = []
    if "total_revenue" in ds_kpis:
        kpis.append(["Total Network Revenue (YTD)", ds_kpis["total_revenue"], "Optimal"])
    if "unique_products" in ds_kpis:
        kpis.append(["Active Pharmaceutical SKUs", ds_kpis["unique_products"], "Optimal"])
    
    kpis.append(["Composite Business Health Score", health, "Optimal" if isinstance(health, (int, float)) and health >= 80 else "Review Required"])
    
    anoms_count = len(anoms)
    kpis.append(["Supply Chain Anomaly Index", anoms_count, "Review Required" if anoms_count > 0 else "Optimal"])

    for row_idx, data_row in enumerate(kpis, 6):
        for col_idx, val in enumerate(data_row, 1):
            cell = ws1.cell(row=row_idx, column=col_idx, value=val)
            cell.font = body_font
            cell.border = thin_border
            if col_idx == 2 and isinstance(val, (int, float)) and val > 1000:
                cell.number_format = "₹#,##0"
                cell.alignment = Alignment(horizontal="right")
            elif col_idx == 3:
                cell.alignment = Alignment(horizontal="center")
                if val == "Optimal":
                    cell.fill = green_fill
                    cell.font = green_font
                else:
                    cell.fill = amber_fill
                    cell.font = amber_font

    # Sheet 2: SKU Performance & Revenue
    ws2 = wb.create_sheet(title="SKU Revenue Breakdown")
    ws2.views.sheetView[0].showGridLines = True

    ws2.merge_cells("A1:D2")
    t2 = ws2["A1"]
    t2.value = "  PHARMACEUTICAL SKU REVENUE BREAKDOWN"
    t2.font = title_font
    t2.fill = navy_fill
    t2.alignment = Alignment(vertical="center")

    headers_sku = ["SKU Code", "Product Name", "Total Revenue (INR)", "Status"]
    ws2.append([])
    ws2.append(headers_sku)

    for col_num in range(1, len(headers_sku) + 1):
        c = ws2.cell(row=4, column=col_num)
        c.font = header_font
        c.fill = indigo_fill
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = thin_border

    skus = []
    top_products = summary.get("sales", {}).get("top_10_products", [])
    if top_products:
        for p in top_products:
            name = str(p.get("product", ""))
            sku = "".join([c.upper() for c in name if c.isalnum()])[:8]
            rev = p.get("revenue", 0)
            skus.append([sku, name, rev, "Optimal"])

    for r_idx, s_data in enumerate(skus, 5):
        bg_fill = light_grey_fill if r_idx % 2 == 0 else PatternFill(fill_type=None)
        for c_idx, val in enumerate(s_data, 1):
            cell = ws2.cell(row=r_idx, column=c_idx, value=val)
            cell.font = body_font
            cell.border = thin_border
            cell.fill = bg_fill
            if c_idx == 3:
                cell.number_format = "₹#,##0.00"
                cell.alignment = Alignment(horizontal="right")
            elif c_idx == 4:
                cell.alignment = Alignment(horizontal="center")
                if val == "Optimal":
                    cell.fill = green_fill
                    cell.font = green_font
                else:
                    cell.fill = red_fill

    if skus:
        tot_row = len(skus) + 5
        for c in range(1, len(headers_sku) + 1):
            t_cell = ws2.cell(row=tot_row, column=c)
            t_cell.fill = soft_indigo_fill
            t_cell.border = thin_border

        c_label = ws2.cell(row=tot_row, column=1, value="TOTAL REVENUE")
        c_label.font = bold_font
        c_val = ws2.cell(row=tot_row, column=3, value=f"=SUM(C5:C{tot_row - 1})")
        c_val.font = bold_font
        c_val.number_format = "₹#,##0.00"

    # Sheet 3: Regional Distribution Matrix
    ws3 = wb.create_sheet(title="Regional Distribution Matrix")
    ws3.views.sheetView[0].showGridLines = True

    ws3.merge_cells("A1:C2")
    t3 = ws3["A1"]
    t3.value = "  REGIONAL DISTRIBUTOR PERFORMANCE"
    t3.font = title_font
    t3.fill = navy_fill
    t3.alignment = Alignment(vertical="center")

    headers_reg = ["Region / Hub", "Revenue Contribution (INR)", "Status"]
    ws3.append([])
    ws3.append(headers_reg)

    for col_num in range(1, len(headers_reg) + 1):
        c = ws3.cell(row=4, column=col_num)
        c.font = header_font
        c.fill = indigo_fill
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = thin_border

    regions = []
    top_regions = summary.get("sales", {}).get("by_region", [])
    if top_regions:
        for r in top_regions:
            reg_name = r.get("name", "Unknown Region")
            rev = r.get("revenue", 0)
            regions.append([reg_name, rev, "Optimal"])

    for r_idx, reg_data in enumerate(regions, 5):
        for c_idx, val in enumerate(reg_data, 1):
            cell = ws3.cell(row=r_idx, column=c_idx, value=val)
            cell.font = body_font
            cell.border = thin_border
            if c_idx == 2:
                cell.number_format = "₹#,##0.00"
                cell.alignment = Alignment(horizontal="right")
            elif c_idx == 3:
                cell.alignment = Alignment(horizontal="center")
                if val == "Optimal":
                    cell.fill = green_fill
                    cell.font = green_font
                else:
                    cell.fill = amber_fill
                    cell.font = amber_font

    for ws in [ws1, ws2, ws3]:
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                if cell.row <= 2 and cell.coordinate != f"{col_letter}3":
                    continue
                if cell.value is not None:
                    val_str = str(cell.value)
                    if not str(cell.value).startswith("="):
                        max_len = max(max_len, len(val_str))
            ws.column_dimensions[col_letter].width = max(max_len + 4, 15)

    out = io.BytesIO()
    wb.save(out)
    out.seek(0)
    return out.getvalue()


class ExecutivePDF(FPDF):
    def __init__(self, title_text: str):
        super().__init__()
        self.doc_title = title_text

    def header(self):
        self.set_fill_color(15, 23, 42)
        self.rect(10, 10, 190, 16, style="F")
        self.set_font("Helvetica", "B", 14)
        self.set_text_color(255, 255, 255)
        self.set_xy(14, 13)
        self.cell(100, 10, f"ASCLEPIUS | {self.doc_title.upper()}", border=0, align="L")
        self.set_font("Helvetica", "I", 9)
        self.set_xy(130, 13)
        self.cell(64, 10, f"Generated: {datetime.now(UTC).strftime('%b %d, %Y')}", border=0, align="R")
        self.set_text_color(0, 0, 0)
        self.ln(20)

    def footer(self):
        self.set_y(-20)
        self.set_font("Helvetica", "I", 8)
        self.set_text_color(100, 116, 139)
        self.cell(0, 10, "Asclepius - Immutable Agent Audit Trail | Confidential Executive Briefing", align="L")
        self.cell(0, 10, f"Page {self.page_no()}", align="R")


def generate_pdf_report(title: str, dataset=None) -> bytes:
    if not dataset or not dataset.analysis_summary:
        raise ValueError("Cannot generate report: No dataset or analysis summary provided.")
    
    summary = dataset.analysis_summary
    clean_title = title.replace("_", " ").replace("PDF", "").replace("XLSX", "").strip()
    pdf = ExecutivePDF(title_text=clean_title or "Executive Report")
    pdf.set_auto_page_break(auto=True, margin=25)
    pdf.add_page()
    
    ds_kpis = summary.get("kpis", {})
    rev = ds_kpis.get("total_revenue", 0)
    rev_str = f"INR {rev / 10000000:.2f} Cr" if rev > 10000000 else f"INR {rev:,.2f}"
    active_skus = str(ds_kpis.get("unique_products", "N/A"))
    anom_list = summary.get("anomalies", [])
    anomalies_count = len(anom_list)
    health_score = f"{dataset.business_health_score:.1f}" if dataset.business_health_score else "N/A"

    # Executive KPI Summary Box
    pdf.set_fill_color(238, 242, 255)
    pdf.set_draw_color(99, 102, 241)
    pdf.set_line_width(0.5)
    pdf.rect(10, 32, 190, 32, style="DF")

    pdf.set_xy(15, 35)
    pdf.set_font("Helvetica", "B", 11)
    pdf.set_text_color(55, 48, 163)
    pdf.cell(180, 6, "EXECUTIVE OPERATIONAL HIGHLIGHTS & COMPOSITE HEALTH", border=0, align="L")

    pdf.set_xy(15, 43)
    pdf.set_font("Helvetica", "B", 10)
    pdf.set_text_color(15, 23, 42)
    pdf.cell(45, 6, f"Health Score: {health_score} / 100", border=0, align="L")
    pdf.cell(45, 6, f"Revenue YTD: {rev_str}", border=0, align="L")
    pdf.cell(45, 6, f"Active SKUs: {active_skus} Lines", border=0, align="L")
    pdf.cell(45, 6, f"Anomalies: {anomalies_count} Active Alerts", border=0, align="L")

    pdf.set_xy(15, 51)
    pdf.set_font("Helvetica", "", 9)
    pdf.set_text_color(71, 85, 105)
    sub_msg = "* Asclepius extracted insights exclusively from the provided dataset."
    pdf.cell(180, 5, sub_msg, border=0)
    pdf.cell(0, 5, "", new_x="LMARGIN", new_y="NEXT")
    pdf.ln(15)

    # Section 1: AI Insights
    pdf.set_font("Helvetica", "B", 13)
    pdf.set_text_color(15, 23, 42)
    pdf.cell(0, 8, "1. Strategic AI Synthesis & Contextual Narrative", new_x="LMARGIN", new_y="NEXT")
    pdf.set_draw_color(203, 213, 225)
    pdf.set_line_width(0.2)
    pdf.line(10, pdf.get_y(), 200, pdf.get_y())
    pdf.ln(4)

    pdf.set_font("Helvetica", "", 10.5)
    pdf.set_text_color(51, 65, 85)
    
    # Use Groq if available
    ai_text = get_ai_narrative(clean_title, summary)
    
    if ai_text:
        paragraphs = [p.strip() for p in ai_text.split("\n\n") if p.strip()]
        for p in paragraphs:
            pdf.multi_cell(0, 5.5, p)
            pdf.ln(3)
    else:
        # Fallback text if groq fails or not configured
        top_regs = summary.get("sales", {}).get("by_region", [])
        if len(top_regs) >= 2:
            r1 = top_regs[0].get("name", "Unknown")
            r2 = top_regs[1].get("name", "Unknown")
            region_str = f"in {r1} and {r2} hubs"
        elif len(top_regs) == 1:
            region_str = f"in the {top_regs[0].get('name', 'Unknown')} hub"
        else:
            region_str = "across the network"

        p1 = (
            "Asclepius autonomous agents conducted diagnostics across distribution networks. "
            f"Trailing performance demonstrates volume momentum {region_str}, "
            f"generating {rev_str} across {active_skus} active commercial SKUs."
        )
        pdf.multi_cell(0, 5.5, p1)
    
    pdf.ln(8)

    t_lower = clean_title.lower()
    
    # If Sales Report, prioritize Regions
    if "sales" in t_lower:
        pdf.set_font("Helvetica", "B", 13)
        pdf.set_text_color(15, 23, 42)
        pdf.cell(0, 8, "2. Top Regional Performance Breakdown", new_x="LMARGIN", new_y="NEXT")
        pdf.line(10, pdf.get_y(), 200, pdf.get_y())
        pdf.ln(4)
        
        top_regions = summary.get("sales", {}).get("by_region", [])
        if not top_regions:
            pdf.set_font("Helvetica", "I", 10)
            pdf.set_text_color(100, 116, 139)
            pdf.multi_cell(0, 5.5, "No regional sales data available.")
        else:
            cols_w = [90, 60]
            pdf.set_font("Helvetica", "B", 9.5)
            pdf.set_fill_color(55, 48, 163)
            pdf.set_text_color(255, 255, 255)
            pdf.cell(cols_w[0], 8, "Region Name", border=1, fill=True, align="L")
            pdf.cell(cols_w[1], 8, "Revenue (INR)", border=1, fill=True, align="R")
            pdf.cell(0, 8, "", new_x="LMARGIN", new_y="NEXT")
            
            pdf.set_font("Helvetica", "", 9)
            for r_idx, reg in enumerate(top_regions[:10]):
                fill = r_idx % 2 == 1
                pdf.set_fill_color(248, 250, 252)
                pdf.set_text_color(51, 65, 85)
                
                rev_val = reg.get("revenue", 0)
                rev = f"{rev_val:,.0f}"
                pdf.cell(cols_w[0], 7.5, reg.get("name", "Unknown"), border=1, fill=fill, align="L")
                pdf.cell(cols_w[1], 7.5, rev, border=1, fill=fill, align="R")
                pdf.cell(0, 7.5, "", new_x="LMARGIN", new_y="NEXT")
        pdf.ln(8)
    
    elif "inventory" in t_lower:
        # If Inventory Report, show Fast Movers
        pdf.set_font("Helvetica", "B", 13)
        pdf.set_text_color(15, 23, 42)
        pdf.cell(0, 8, "2. Fast Moving Inventory Items", new_x="LMARGIN", new_y="NEXT")
        pdf.line(10, pdf.get_y(), 200, pdf.get_y())
        pdf.ln(4)
        
        fast_movers = summary.get("inventory", {}).get("fast_movers", [])
        if not fast_movers:
            pdf.set_font("Helvetica", "I", 10)
            pdf.set_text_color(100, 116, 139)
            pdf.multi_cell(0, 5.5, "No fast moving inventory data available.")
        else:
            cols_w = [80, 40, 30]
            pdf.set_font("Helvetica", "B", 9.5)
            pdf.set_fill_color(55, 48, 163)
            pdf.set_text_color(255, 255, 255)
            pdf.cell(cols_w[0], 8, "Product Name", border=1, fill=True, align="L")
            pdf.cell(cols_w[1], 8, "Velocity (Days)", border=1, fill=True, align="C")
            pdf.cell(cols_w[2], 8, "Turnover", border=1, fill=True, align="R")
            pdf.cell(0, 8, "", new_x="LMARGIN", new_y="NEXT")
            
            pdf.set_font("Helvetica", "", 9)
            for r_idx, item in enumerate(fast_movers[:10]):
                fill = r_idx % 2 == 1
                pdf.set_fill_color(248, 250, 252)
                pdf.set_text_color(51, 65, 85)
                
                pdf.cell(cols_w[0], 7.5, str(item.get("product", ""))[:35], border=1, fill=fill, align="L")
                pdf.cell(cols_w[1], 7.5, str(item.get("avg_days_to_sell", "-")), border=1, fill=fill, align="C")
                pdf.cell(cols_w[2], 7.5, f"{item.get('turnover_rate', 0):.1f}x", border=1, fill=fill, align="R")
                pdf.cell(0, 7.5, "", new_x="LMARGIN", new_y="NEXT")
        pdf.ln(8)
    else:
        # Default Section 2: Anomalies
        pdf.set_font("Helvetica", "B", 13)
        pdf.set_text_color(15, 23, 42)
        pdf.cell(0, 8, "2. High-Priority Supply Chain Anomalies & Action Items", new_x="LMARGIN", new_y="NEXT")
        pdf.line(10, pdf.get_y(), 200, pdf.get_y())
        pdf.ln(4)

        if not anom_list:
            pdf.set_font("Helvetica", "I", 10)
            pdf.set_text_color(100, 116, 139)
            pdf.multi_cell(0, 5.5, "No critical anomalies or supply chain alerts detected in this period.")
            pdf.ln(8)
        else:
            for idx, a in enumerate(anom_list[:2]):
                severity = a.get('severity', 'anomaly').lower()
                title_text = f"[{severity.upper()}] - {a.get('product', 'Network Alert')}"
                desc_text = a.get('description', 'Detected anomaly in supply chain.')
                rec_text = a.get('recommendation', '')
                if rec_text:
                    desc_text += f" Recommendation: {rec_text}"

                if severity == "critical":
                    pdf.set_fill_color(254, 242, 242)
                    pdf.set_draw_color(239, 68, 68)
                    text_col = (153, 27, 27)
                else:
                    pdf.set_fill_color(254, 243, 199)
                    pdf.set_draw_color(245, 158, 11)
                    text_col = (146, 64, 14)

                pdf.rect(10, pdf.get_y(), 190, 24, style="DF")
                pdf.set_xy(14, pdf.get_y() + 2)
                pdf.set_font("Helvetica", "B", 10)
                pdf.set_text_color(*text_col)
                pdf.cell(0, 6, title_text, new_x="LMARGIN", new_y="NEXT")
                pdf.set_xy(14, pdf.get_y())
                pdf.set_font("Helvetica", "", 9.5)
                pdf.set_text_color(51, 65, 85)
                pdf.multi_cell(182, 5, desc_text)
                pdf.ln(12)


    # Section 3: SKU Table (always relevant)
    pdf.add_page()
    pdf.set_font("Helvetica", "B", 13)
    pdf.set_text_color(15, 23, 42)
    pdf.cell(0, 8, "3. Top Pharmaceutical SKUs - Revenue Breakdown", new_x="LMARGIN", new_y="NEXT")
    pdf.line(10, pdf.get_y(), 200, pdf.get_y())
    pdf.ln(4)

    cols_w = [35, 90, 35, 30]
    headers_pdf = ["SKU Code", "Product Name", "Revenue (INR)", "Status"]
    pdf.set_font("Helvetica", "B", 9.5)
    pdf.set_fill_color(55, 48, 163)
    pdf.set_text_color(255, 255, 255)

    for idx, h_txt in enumerate(headers_pdf):
        align = "C" if idx in (0, 3) else ("R" if idx == 2 else "L")
        pdf.cell(cols_w[idx], 8, h_txt, border=1, fill=True, align=align)
    pdf.cell(0, 8, "", new_x="LMARGIN", new_y="NEXT")

    top_products = summary.get("sales", {}).get("top_10_products", [])
    if not top_products:
        pdf.set_font("Helvetica", "I", 10)
        pdf.set_text_color(100, 116, 139)
        pdf.cell(190, 10, "No SKU revenue data available.", border=1, align="C", new_x="LMARGIN", new_y="NEXT")
    else:
        pdf.set_font("Helvetica", "", 9)
        pdf.set_draw_color(203, 213, 225)
        for r_idx, p in enumerate(top_products[:15]):
            fill = r_idx % 2 == 1
            pdf.set_fill_color(248, 250, 252)
            
            name = str(p.get("product", ""))
            sku = "".join([c.upper() for c in name if c.isalnum()])[:8]
            rev_val = p.get("revenue", 0)
            rev = f"{rev_val:,.0f}"
            status = "Optimal"

            row_data = [sku, name[:40], rev, status]
            
            for idx, val in enumerate(row_data):
                align = "C" if idx in (0, 3) else ("R" if idx == 2 else "L")
                if idx == 3:
                    pdf.set_text_color(22, 101, 52)
                else:
                    pdf.set_text_color(51, 65, 85)
                pdf.cell(cols_w[idx], 7.5, str(val), border=1, fill=fill, align=align)
            pdf.cell(0, 7.5, "", new_x="LMARGIN", new_y="NEXT")

    pdf_bytes = bytes(pdf.output())
    return pdf_bytes
